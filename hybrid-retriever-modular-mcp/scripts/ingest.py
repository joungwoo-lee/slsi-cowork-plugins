"""Local document ingest: copy file, extract text, chunk, index SQLite and Qdrant."""
from __future__ import annotations

import hashlib
import shutil
from pathlib import Path
from typing import Any

from .config import Config
from .embedding_client import EmbeddingClient
from . import storage


def read_text(path: Path, max_chars: int) -> str:
    suffix = path.suffix.lower()
    if suffix == ".pdf":
        from pypdf import PdfReader

        reader = PdfReader(path)
        text = "\n\n".join((page.extract_text() or "") for page in reader.pages)
    elif suffix in (".doc", ".docx"):
        from docx import Document

        doc = Document(path)
        text = "\n\n".join(p.text for p in doc.paragraphs if p.text)
    elif suffix in (".xls", ".xlsx"):
        import openpyxl

        wb = openpyxl.load_workbook(path, read_only=True, data_only=True)
        lines: list[str] = []
        for sheet in wb.worksheets:
            lines.append(f"# Sheet: {sheet.title}")
            for row in sheet.iter_rows(values_only=True):
                line = "\t".join(str(cell) if cell is not None else "" for cell in row).strip()
                if line:
                    lines.append(line)
        text = "\n".join(lines)
    else:
        data = path.read_bytes()
        try:
            text = data.decode("utf-8")
        except UnicodeDecodeError:
            try:
                import chardet

                encoding = chardet.detect(data).get("encoding") or "utf-8"
                text = data.decode(encoding, errors="replace")
            except Exception:
                text = data.decode("utf-8", errors="replace")
    return text[:max_chars]


def chunk_text(text: str, chunk_chars: int, overlap: int) -> list[str]:
    text = text.strip()
    if not text:
        return []
    chunks: list[str] = []
    step = max(1, chunk_chars - max(0, overlap))
    for start in range(0, len(text), step):
        chunk = text[start : start + chunk_chars].strip()
        if chunk:
            chunks.append(chunk)
        if start + chunk_chars >= len(text):
            break
    return chunks


def _flat_records(chunks: list[str]) -> list[dict[str, Any]]:
    return [
        {
            "child_content": text,
            "original_child_content": text,
            "parent_content": text,
            "parent_id": 0,
            "child_id": idx,
            "global_position": idx,
            "is_hierarchical": False,
            "is_contextual": False,
        }
        for idx, text in enumerate(chunks)
    ]


def hierarchical_records(cfg: Config, text: str, parent_mode: str = "normal") -> list[dict[str, Any]]:
    if parent_mode == "full":
        child_chunks = chunk_text(text, cfg.ingest.child_chunk_chars, cfg.ingest.child_chunk_overlap)
        parents = [(0, text, child_chunks)]
    else:
        parent_chunks = chunk_text(text, cfg.ingest.parent_chunk_chars, cfg.ingest.parent_chunk_overlap)
        parents = [
            (parent_id, parent, chunk_text(parent, cfg.ingest.child_chunk_chars, cfg.ingest.child_chunk_overlap))
            for parent_id, parent in enumerate(parent_chunks)
        ]
    records: list[dict[str, Any]] = []
    for parent_id, parent, child_chunks in parents:
        for child_id, child in enumerate(child_chunks):
            records.append(
                {
                    "child_content": child,
                    "original_child_content": child,
                    "parent_content": parent,
                    "parent_id": parent_id,
                    "child_id": child_id,
                    "global_position": len(records),
                    "is_hierarchical": True,
                    "is_contextual": False,
                }
            )
    return records


def document_id_for(dataset_id: str, path: Path) -> str:
    stat = path.stat()
    raw = f"{dataset_id}|{path.resolve()}|{stat.st_size}|{stat.st_mtime_ns}"
    return hashlib.sha1(raw.encode("utf-8", errors="replace")).hexdigest()[:20]


def upload_document(
    cfg: Config,
    dataset_id: str,
    file_path: str,
    *,
    skip_embedding: bool = False,
    use_hierarchical: str | bool | None = None,
    metadata: dict | None = None,
) -> dict:
    path = Path(file_path).expanduser()
    if not path.is_file():
        raise FileNotFoundError(f"file not found: {path}")
    text = read_text(path, cfg.ingest.max_file_chars)
    if str(use_hierarchical).lower() in ("true", "full"):
        chunk_records = hierarchical_records(cfg, text, parent_mode="full" if str(use_hierarchical).lower() == "full" else "normal")
    else:
        chunk_records = _flat_records(chunk_text(text, cfg.ingest.chunk_chars, cfg.ingest.chunk_overlap))
    if not chunk_records:
        raise ValueError(f"no text content extracted: {path}")

    doc_id = document_id_for(dataset_id, path)
    doc_dir = cfg.document_dir(dataset_id, doc_id)
    doc_dir.mkdir(parents=True, exist_ok=True)
    stored_source = cfg.source_path(dataset_id, doc_id, path.name)
    stored_content = cfg.content_path(dataset_id, doc_id)
    shutil.copy2(path, stored_source)
    stored_content.write_text(text, encoding="utf-8", errors="replace")

    has_vector = False
    vectors: list[list[float]] = []
    if not skip_embedding and cfg.embedding.api_url and cfg.embedding.dim > 0:
        vectors = EmbeddingClient(cfg.embedding).embed([r["child_content"] for r in chunk_records])
        has_vector = True

    with storage.sqlite_session(cfg) as conn:
        storage.ensure_dataset(conn, dataset_id)
        storage.upsert_document(
            conn,
            dataset_id=dataset_id,
            document_id=doc_id,
            name=path.name,
            source_path=str(stored_source),
            content_path=str(stored_content),
            size_bytes=path.stat().st_size,
            chunks=chunk_records,
            has_vector=has_vector,
            metadata=metadata,
        )

    if has_vector:
        client = storage.open_qdrant(cfg)
        storage.ensure_collection(client, cfg)
        storage.upsert_vectors(
            client,
            cfg,
            [
                (
                    f"{doc_id}:{pos}",
                    vector,
                    {
                        "dataset_id": dataset_id,
                        "document_id": doc_id,
                        "document_name": path.name,
                        "position": pos,
                        "parent_content": chunk_records[pos].get("parent_content", ""),
                        "is_hierarchical": bool(chunk_records[pos].get("is_hierarchical", False)),
                        "metadata": metadata or {},
                    },
                )
                for pos, vector in enumerate(vectors)
            ],
        )

    return {
        "dataset_id": dataset_id,
        "document_id": doc_id,
        "name": path.name,
        "chunks_count": len(chunk_records),
        "parent_chunks_count": len({r.get("parent_id") for r in chunk_records if r.get("is_hierarchical")}),
        "is_hierarchical": any(r.get("is_hierarchical") for r in chunk_records),
        "has_vector": has_vector,
        "source_path": str(stored_source),
        "content_path": str(stored_content),
    }
